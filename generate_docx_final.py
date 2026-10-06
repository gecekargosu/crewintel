"""Generate DOCX from the final CV content."""
import os

output_dir = r"C:\Users\isitm\Desktop\Cengiz_Kilic_CV_Yenilenmis"

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)
style.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.15

def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
        run.font.name = 'Calibri'
    return h

def add_para(text, bold=False, italic=False, size=10, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    if color:
        run.font.color.rgb = RGBColor(*color)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run_bold = p.add_run(bold_prefix)
        run_bold.bold = True
        run_bold.font.size = Pt(9.5)
        run_bold.font.name = 'Calibri'
    run = p.add_run(text)
    run.font.size = Pt(9.5)
    run.font.name = 'Calibri'
    p.paragraph_format.space_after = Pt(2)
    return p

# HEADER
add_para("Cengiz K\u0131l\u0131\u00e7", bold=True, size=22, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
add_para("AI & Machine Learning Student \u00b7 Full-Stack Developer \u00b7 AI Automation Specialist", size=10, color=(0x44,0x44,0x44), space_after=4)
add_para("+90 532 327 61 21  |  cengizkilic1980@gmail.com  |  \u0130zmir, T\u00fcrkiye", size=9, color=(0x33,0x33,0x33), space_after=1)
add_para("linkedin.com/in/cengizkilic  |  github.com/cengizkilic", size=9, color=(0x33,0x33,0x33), space_after=1)
add_para("Full-time \u00b7 Remote / Hybrid \u00b7 No travel restrictions", size=9, color=(0x55,0x55,0x55), space_after=8)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# PROFESYONEL \u00d6ZET
add_heading_styled("PROFESYONEL \u00d6ZET", level=2)
add_para("20 y\u0131l\u0131 a\u015f\u0131n profesyonel deneyimini, giri\u015fimcilik ge\u00e7mi\u015fini ve \u00fcr\u00fcn geli\u015ftirme bilgisini AI/ML, software development ve automation alan\u0131na ta\u015f\u0131yan bir geli\u015ftirici. \u0130stanbul Rumeli \u00d6niversitesi Yapay Zeka ve Makine \u00d6\u011frenmesi lisans program\u0131nda e\u011fitimine devam ediyor.", space_after=4)
add_para("CrewIntel adl\u0131 full-stack crew management sistemini s\u0131f\u0131rdan geli\u015ftirerek Python/FastAPI backend, PostgreSQL veritaban\u0131, React frontend, Android mobil uygulama ve Groq LLM entegrasyonu i\u00e7eren tam bir AI uygulamas\u0131 olu\u015fturdu. n8n workflow otomasyonu, prompt engineering, AI agent tasar\u0131m\u0131 ve LLM tabanl\u0131 uygulama geli\u015ftirme konular\u0131nda aktif \u00e7al\u0131\u015fmalar y\u00f6rut\u00fcyor.", space_after=8)

# E\u011e\u0130T\u0130M
add_heading_styled("E\u011e\u0130T\u0130M", level=2)
p = doc.add_paragraph()
run = p.add_run("\u0130stanbul Rumeli \u00d6niversitesi \u2014 M\u00fchendislik Fak\u00fcltesi")
run.bold = True; run.font.size = Pt(10); run.font.name = 'Calibri'
run2 = p.add_run("                                                                                              2025 \u2013 Devam Ediyor (Tahmini: 2029)")
run2.font.size = Pt(9); run2.font.color.rgb = RGBColor(0x55,0x55,0x55); run2.font.name = 'Calibri'
p.paragraph_format.space_after = Pt(1)
add_para("Yapay Zeka ve Makine \u00d6\u011frenmesi \u2014 Lisans Program\u0131", size=9.5, color=(0x33,0x33,0x33), space_after=1)
add_para("B\u00f6l\u00fcm \u00d6\u011frenci Temsilcisi \u00b7 Algoritmalar, Veri Yap\u0131lar\u0131, Machine Learning, Deep Learning, Software Development, AI Applications", size=9, color=(0x55,0x55,0x55), space_after=6)

p = doc.add_paragraph()
run = p.add_run("Karadeniz Teknik \u00d6niversitesi")
run.bold = True; run.font.size = Pt(10); run.font.name = 'Calibri'
run2 = p.add_run("                                                                                              1999 \u2013 2001 \u00b7 Mezun")
run2.font.size = Pt(9); run2.font.color.rgb = RGBColor(0x55,0x55,0x55); run2.font.name = 'Calibri'
p.paragraph_format.space_after = Pt(1)
add_para("\u0130n\u015faat Teknikerli\u011fi \u2014 \u00d6n Lisans", size=9.5, color=(0x33,0x33,0x33), space_after=8)

# TEKN\u0130K YETK\u0130NL\u0130KLER
add_heading_styled("TEKN\u0130K YETK\u0130NL\u0130KLER", level=2)
add_bullet("LLM (Groq / llama-3.3-70b-versatile) \u00b7 Prompt Engineering \u00b7 AI Agents \u00b7 AI Application Development \u00b7 Document Analysis with LLM \u00b7 AI-powered Matching Systems", bold_prefix="AI & Machine Learning: ")
add_bullet("n8n \u00b7 API Integration \u00b7 Workflow Automation \u00b7 AI-assisted Automation \u00b7 Multi-step Prompt Architectures", bold_prefix="AI Automation & Workflow: ")
add_bullet("Python \u00b7 FastAPI \u00b7 SQLAlchemy \u00b7 PostgreSQL \u00b7 Pydantic \u00b7 Alembic \u00b7 REST API \u00b7 JWT Authentication", bold_prefix="Backend & Programming: ")
add_bullet("HTML \u00b7 CSS \u00b7 JavaScript \u00b7 React \u00b7 Vite \u00b7 Android (Java/Kotlin) \u00b7 Retrofit", bold_prefix="Frontend & Mobile: ")
add_bullet("Git \u00b7 GitHub \u00b7 Docker \u00b7 Docker Compose \u00b7 VS Code \u00b7 Linux \u00b7 WSL", bold_prefix="DevOps & Developer Tools: ")
add_bullet("ChatGPT \u00b7 Claude \u00b7 Gemini \u00b7 GitHub Copilot \u00b7 Cursor \u00b7 C (Basic)", bold_prefix="AI Tools & Others: ")

# PROJELER
add_heading_styled("PROJELER", level=2)
add_para("CrewIntel \u2014 Yapay Zeka Destekli M\u00fcrettebat Y\u00f6netim Sistemi", bold=True, size=10.5, space_after=1)
add_para("Python \u00b7 FastAPI \u00b7 PostgreSQL \u00b7 SQLAlchemy \u00b7 React \u00b7 Vite \u00b7 Android \u00b7 Groq LLM \u00b7 Docker Compose", italic=True, size=9, color=(0x55,0x55,0x55), space_after=3)
add_bullet("geli\u015ftirerek 20+ REST API endpoint'i, JWT kimlik do\u011fulama, CORS, rate limiting ve audit logging i\u00e7eren backend mimarisi kurdum.", bold_prefix="Python/FastAPI backend ")
add_bullet("ile crew, document, ship, contract, assignment gibi veri modellerini tasarlad\u0131m ve Alembic ile y\u00f6netim uygulad\u0131m.", bold_prefix="PostgreSQL + SQLAlchemy ORM ")
add_bullet("ile PDF belgelerden personel bilgileri, sertifikalar ve deneyim verilerini \u00e7\u0131kar\u0131an AI document analyzer mod\u00fcl\u00fc geli\u015ftirdim.", bold_prefix="Groq LLM API entegrasyonu ")
add_bullet("\u2014 LLM ve algoritmik skorlama ile personel yeteneklerini, sertifikalar\u0131n\u0131 ve deneyimlerini i\u015f ilanlar\u0131yla e\u015fle\u015ftiren matching motoru tasarlad\u0131m.", bold_prefix="AI-powered crew matcher ")
add_bullet("\u2014 Belgelerde tutars\u0131zl\u0131k ve veri anomali tespiti yapan AI tabanl\u0131 kontrol mod\u00fcl\u00fc olu\u015fturdum.", bold_prefix="Anomaly detector ")
add_bullet("(Java, Retrofit) geli\u015ftirerek admin paneli ve mobil eri\u015fim sa\u011flad\u0131m.", bold_prefix="React + Vite frontend ve Android mobil uygulama ")
add_bullet("ile multi-container deployment (PostgreSQL + Backend) yap\u0131land\u0131rarak ta\u015f\u0131nabilir development environment kurdum.", bold_prefix="Docker Compose ")
add_bullet("E-posta ve WhatsApp bildirim sistemi, maa\u015f takibi, i\u015f ilan\u0131 y\u00f6netimi, belge y\u00fckleme ve dinamik dashboard \u00f6zelliklerini uygulad\u0131m.")
doc.add_paragraph().paragraph_format.space_after = Pt(4)

add_para("Social Media Downloader \u2014 Sosyal Medya Video \u0130ndirici API", bold=True, size=10.5, space_after=1)
add_para("Python \u00b7 FastAPI \u00b7 yt-dlp \u00b7 Docker", italic=True, size=9, color=(0x55,0x55,0x55), space_after=3)
add_bullet("YouTube, Instagram, TikTok, Facebook ve Pinterest videolar\u0131n\u0131 indiren FastAPI tabanl\u0131 REST API geli\u015ftirdim.")
add_bullet("entegrasyonu ile \u00e7oklu platform deste\u011fi, video metadata \u00e7\u0131karma ve dosya y\u00f6netimi sa\u011flad\u0131m.", bold_prefix="yt-dlp ")
add_bullet("containerization ile ta\u015f\u0131nabilir deployment mimarisi kurdum.", bold_prefix="Docker ")

# G\u0130R\u0130\u015e\u0130MC\u0130L\u0130K
add_heading_styled("G\u0130R\u0130\u015e\u0130MC\u0130L\u0130K DENEY\u0130M\u0130", level=2)
p = doc.add_paragraph()
run = p.add_run("NGD \u2014 New Generation Design \u00b7 Kurucu & \u00dcr\u00fcn Geli\u015ftirici")
run.bold = True; run.font.size = Pt(10); run.font.name = 'Calibri'
run2 = p.add_run("                                                                                              2016 \u2013 2022")
run2.font.size = Pt(9); run2.font.color.rgb = RGBColor(0x55,0x55,0x55); run2.font.name = 'Calibri'
p.paragraph_format.space_after = Pt(1)
add_para("Termal y\u00f6netim, teknik tekstil ve savunma sanayii \u2014 Ar-Ge odakl\u0131 \u00fcr\u00fcn geli\u015ftirme", size=9.5, color=(0x44,0x44,0x44), space_after=3)
add_bullet("\u015eirketi s\u0131f\u0131rdan kurdu; m\u00fc\u015fteri problemini tan\u0131mlad\u0131, \u00e7\u00f6z\u00fcm hipotezi olu\u015fturdu, prototip geli\u015ftirdi ve MVP'yi piyasaya \u00e7\u0131kar\u0131d\u0131 \u2014 start-up s\u00fcre\u00e7lerini u\u00e7tan uca y\u00f6netti.")
add_bullet("M\u00fc\u015fteri ihtiya\u00e7 analizi, \u00fcr\u00fcn-pazar uyumu ara\u015ft\u0131rmas\u0131 ve \u00fcr\u00fcn do\u011frulama s\u00fcre\u00e7lerini y\u00f6rtt\u00fc; B2B m\u00fc\u015fteriler i\u00e7in \u00f6zel \u00e7\u00f6z\u00fcmler geli\u015ftirdi.")
add_bullet("Kurumsal web sitesi, dijital pazarlama, \u00fcr\u00fcn lansman\u0131 ve e-ticaret s\u00fcre\u00e7lerini y\u00f6netti.")
add_bullet("Savunma sanayii i\u00e7in \u00f6zel termal sistemler geli\u015ftirdi; Milli Savunma Bakanl\u0131\u011f\u0131 ve Silahl\u0131 Kuvvetler'e teknik sunum yapt\u0131.")
add_bullet("Almanya, Fransa, \u0130sve\u00e7, ABD ve Rusya dahil uluslararas\u0131 tedarik\u00e7i ve ihracat s\u00fcre\u00e7lerini koordine etti.")

# MESLEK\u0130 GE\u00c7M\u0130\u015e
add_heading_styled("MESLEK\u0130 GE\u00c7M\u0130\u015e", level=2)
add_para("Yurt i\u00e7i ve d\u0131\u015f\u0131nda yakla\u015f\u0131k 20 y\u0131l boyunca teknik operasyon, proje koordinasyonu, ekip y\u00f6netimi ve kalite kontrol alanlar\u0131nda g\u00f6rev ald\u0131:", size=9.5, color=(0x33,0x33,0x33), space_after=3)
add_bullet("Farkl\u0131 \u00fclkelerde ve disiplinlerde saha deneyimi kazanarak de\u011fi\u015fen \u015fartlara uyum sa\u011flama ve sonu\u00e7 odakl\u0131 hareket etme becerilerini geli\u015ftirdi.")
add_bullet("Teknik operasyon ve proje y\u00f6netiminde y\u00fcksek sorumluluk gerektiren g\u00f6revlerde analitik bak\u0131\u015f a\u00e7\u0131s\u0131 ve problem \u00e7\u00f6zme yetkinli\u011fini ortaya koydu.")
add_bullet("Uluslararas\u0131 projelerde m\u00fc\u015fteri ve tedarik\u00e7i ileti\u015fimi, kalite kontrol ve ekip koordinasyonu s\u00fcre\u00e7lerini y\u00f6rtt\u00fc.")
add_bullet("Bu deneyimler, yaz\u0131l\u0131m geli\u015ftirme s\u00fcrecinde Agile thinking, technical problem solving, cross-functional collaboration ve project management gibi transferable skills olarak do\u011frudan katk\u0131 sa\u011flamaktad\u0131r.")

# N\u0130TEL\u0130KLER
add_heading_styled("N\u0130TEL\u0130KLER", level=2)
add_para("Analitik d\u00fc\u015f\u00fcnme \u00b7 Problem \u00e7\u00f6zme \u00b7 Giri\u015fimci bak\u0131\u015f a\u00e7\u0131s\u0131 \u00b7 Tak\u0131m \u00e7al\u0131\u015fmas\u0131 \u00b7 Sonu\u00e7 odakl\u0131 \u00b7 H\u0131zl\u0131 \u00f6\u011frenme", size=9.5, color=(0x33,0x33,0x33), space_after=6)

# YABANCI D\u0130L
add_heading_styled("YABANCI D\u0130L", level=2)
add_para("T\u00fcrk\u00e7e \u2014 Ana dil", bold=True, size=9.5, space_after=1)
add_para("\u0130ngilizce \u2014 B2 / Upper-Intermediate \u00b7 Teknik d\u00f6k\u00fcman okuma, yazma ve temel konu\u015fma", size=9, color=(0x55,0x55,0x55), space_after=6)

# KAR\u0130YER HEDEF\u0130
add_heading_styled("KAR\u0130YER HEDEF\u0130", level=2)
add_para("AI/ML, LLM uygulamalar\u0131, AI automation ve backend development alanlar\u0131nda uzmanla\u015f\u0131arak ger\u00e7ek problemlere y\u00fnelik yaz\u0131l\u0131m \u00e7\u00f6z\u00fcmleri geli\u015ftirmeyi hedefliyorum. Full-stack AI proje deneyimi, AI/ML lisans e\u011fitimi ve 20+ y\u0131ll\u0131k profesyonel/problem-solving ge\u00e7mi\u015fimi birle\u015ftirerek Junior AI Developer, AI Automation Developer ve Junior Backend rollerinde de\u011fer \u00fcretmeye odaklan\u0131yorum.", size=9.5, space_after=8)

# REFERANSLAR
add_heading_styled("REFERANSLAR", level=2)
add_para("Talep edilmesi halinde profesyonel referanslar\u0131m payla\u015f\u0131lacakt\u0131r.", size=9, color=(0x55,0x55,0x55), space_after=4)

# FOOTER
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("CV Updated: September 2026")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(0x99,0x99,0x99)
run.font.name = 'Calibri'

docx_path = os.path.join(output_dir, "cengiz_kilic_cv_final.docx")
doc.save(docx_path)
print(f"DOCX generated: {docx_path}")

# Verify
for f in os.listdir(output_dir):
    fpath = os.path.join(output_dir, f)
    size = os.path.getsize(fpath)
    print(f"  {f} ({size:,} bytes)")
