# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

style = doc.styles['Normal']
style.font.name = 'Segoe UI'
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(2)
style.paragraph_format.line_spacing = 1.15

def add_section_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.name = 'Segoe UI'
    run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
    border = doc.add_paragraph()
    border.paragraph_format.space_after = Pt(4)
    run2 = border.add_run('_' * 95)
    run2.font.color.rgb = RGBColor(0xcc, 0xcc, 0xcc)
    run2.font.size = Pt(5)
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.5)
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(9.8)
        run.font.name = 'Segoe UI'
        run = p.add_run(text)
        run.font.size = Pt(9.8)
        run.font.name = 'Segoe UI'
    else:
        run = p.add_run(text)
        run.font.size = Pt(9.8)
        run.font.name = 'Segoe UI'
    return p

def add_simple_para(text, size=10.5, color=None, bold=False, italic=False, space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = 'Segoe UI'
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    return p

# ═══ HEADER ═══
h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = h.add_run('Cengiz Kılıç')
run.bold = True
run.font.size = Pt(22)
run.font.name = 'Segoe UI'
h.paragraph_format.space_after = Pt(2)

add_simple_para('Yapay Zeka & Makine Öğrenmesi Lisans Öğrencisi · AI Otomasyon Geliştirici · Junior Software Developer', 10.5, (0x44,0x44,0x44), space_after=4)
add_simple_para('+90 532 327 61 21  |  cengizkilic1980@gmail.com  |  İzmir, Türkiye  |  linkedin.com/in/cengizkilic', 9.5, (0x33,0x33,0x33), space_after=2)
add_simple_para('Tam zamanlı · Uzaktan / Hibrit · Seyahat engeli yok', 9.5, (0x55,0x55,0x55), space_after=6)

div = doc.add_paragraph()
div.paragraph_format.space_after = Pt(8)
run = div.add_run('_' * 95)
run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
run.font.size = Pt(6)

# ═══ PROFESYONEL ÖZET ═══
add_section_title('Profesyonel Özet')

for t in [
    '20 yılı aşkın profesyonel deneyimini yapay zeka, yazılım geliştirme ve otomasyon alanına taşıyan bir geliştirici adayı. İstanbul Rumeli Üniversitesi Yapay Zeka ve Makine Öğrenmesi lisans programında eğitimine devam ediyor.',
    'CrewIntel adlı tam yığın (full-stack) mürettebat yönetim sistemi geliştirerek Python FastAPI backend, PostgreSQL veritabanı, React frontend ve Android mobil uygulama oluşturdum. Sistemde Groq LLM entegrasyonu ile belge analizi ve yapay zeka destekli personel-iş eşleştirme motoru tasarladım. Docker ile deployment gerçekleştirdim.',
    'n8n workflow otomasyonu, prompt engineering, AI agent tasarımı ve LLM uygulama geliştirme konularında aktif çalışmalar yürütüyorum. Girişimcilik geçmişim — ürün geliştirme, MVP süreçleri, müşteri analizi ve uluslararası iş geliştirme — yazılım geliştirme sürecine güçlü bir iş disiplini ve problem çözme yetkinliği katıyor.',
]:
    add_simple_para(t, 10.5, space_after=4)

# ═══ EĞİTİM ═══
add_section_title('Eğitim')

p = doc.add_paragraph()
run = p.add_run('İstanbul Rumeli Üniversitesi — Mühendislik Fakültesi')
run.bold = True; run.font.size = Pt(10.5); run.font.name = 'Segoe UI'
run2 = p.add_run('\t2025 – Devam Ediyor (Tahmini: 2029)')
run2.font.size = Pt(9.5); run2.font.color.rgb = RGBColor(0x55,0x55,0x55)
p.paragraph_format.space_after = Pt(1)

add_simple_para('Yapay Zeka ve Makine Öğrenmesi — Lisans Programı', 10, (0x33,0x33,0x33), space_after=1)
add_simple_para('Bölüm Öğrenci Temsilcisi · Algoritmalar, veri yapıları, makine öğrenmesi, derin öğrenme, yazılım geliştirme temelleri ve yapay zeka uygulamaları', 9.5, (0x55,0x55,0x55), space_after=6)

p = doc.add_paragraph()
run = p.add_run('Karadeniz Teknik Üniversitesi')
run.bold = True; run.font.size = Pt(10.5); run.font.name = 'Segoe UI'
run2 = p.add_run('\t1999 – 2001 · Mezun')
run2.font.size = Pt(9.5); run2.font.color.rgb = RGBColor(0x55,0x55,0x55)
p.paragraph_format.space_after = Pt(1)

add_simple_para('İnşaat Teknikerliği — Ön Lisans', 10, (0x33,0x33,0x33), space_after=6)

# ═══ TEKNİK YETKİNLİKLER ═══
add_section_title('Teknik Yetkinlikler')

for label, items in [
    ('AI & Machine Learning', 'LLM (Groq / llama-3.3-70b) · Prompt Engineering · AI Agents · AI Application Development · Document Analysis with LLM · AI-powered Matching Systems'),
    ('AI Automation & Workflow', 'n8n · API Integration · Workflow Automation · AI-assisted Automation · Multi-step Prompt Architectures'),
    ('Backend & Programming', 'Python · FastAPI · SQLAlchemy · PostgreSQL · REST API · JWT Authentication · Pydantic · Alembic'),
    ('Frontend & Mobile', 'HTML · CSS · JavaScript · React · Android (Java/Kotlin) · Retrofit'),
    ('Developer Tools', 'Git · GitHub · VS Code · Docker · Docker Compose · Linux · WSL'),
    ('AI Tools & Others', 'ChatGPT · Claude · Gemini · GitHub Copilot · Cursor · C (Basic) · Figma · LibreOffice'),
]:
    p = doc.add_paragraph()
    run = p.add_run(label + ': ')
    run.bold = True; run.font.size = Pt(10); run.font.name = 'Segoe UI'
    run2 = p.add_run(items)
    run2.font.size = Pt(9.8); run2.font.name = 'Segoe UI'
    p.paragraph_format.space_after = Pt(2)

# ═══ PROJELER ═══
add_section_title('Projeler')

# CrewIntel
p = doc.add_paragraph()
run = p.add_run('CrewIntel — Yapay Zeka Destekli Mürettebat Yönetim Sistemi')
run.bold = True; run.font.size = Pt(10.5)
p.paragraph_format.space_after = Pt(1)

add_simple_para('Tam yığın (full-stack) crew management platformu · Python FastAPI + PostgreSQL + React + Android + LLM', 9.8, (0x44,0x44,0x44), italic=True, space_after=2)

for bp, np in [
    ('Python FastAPI', ' backend geliştirerek 20+ API endpoint\'i, JWT kimlik doğrulama, CORS, rate limiting ve audit logging mimarisi kurdum.'),
    ('PostgreSQL', ' veritabanı ile SQLAlchemy ORM kullanarak crew, document, ship, contract, assignment veri modelleri oluşturdum.'),
    ('Groq LLM API entegrasyonu', ' ile PDF belgelerden personel bilgileri, sertifikalar ve deneyim çıkaran AI document analyzer modülü geliştirdim.'),
    ('AI-powered crew matcher', ' — LLM ve algoritmik skorlama ile personel yeteneklerini, sertifikalarını ve deneyimlerini iş ilanlarıyla eşleştiren motor tasarladım.'),
    ('Anomaly detector', ' — Belgelerde tutarsızlık tespiti yapan AI modülü oluşturdum.'),
    ('React + Vite frontend', ' ve Android mobil uygulama (Java) geliştirerek admin paneli ve mobil erişim sağladım.'),
    ('Docker Compose', ' ile multi-container deployment (PostgreSQL + Backend) yapılandırdım.'),
    (None, 'E-posta ve WhatsApp bildirim sistemi, maaş takibi, iş ilanı yönetimi, belge yükleme ve dinamik dashboard özelliklerini uyguladım.'),
]:
    add_bullet(np, bp)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Social Downloader
p = doc.add_paragraph()
run = p.add_run('Social Media Downloader — Sosyal Medya Video İndirici API')
run.bold = True; run.font.size = Pt(10.5)
p.paragraph_format.space_after = Pt(1)

add_simple_para('FastAPI backend · yt-dlp · Docker', 9.8, (0x44,0x44,0x44), italic=True, space_after=2)

add_bullet(' YouTube, Instagram, TikTok, Facebook ve Pinterest videolarını indiren FastAPI tabanlı REST API geliştirdim.')
add_bullet('yt-dlp entegrasyonu ile çoklu platform desteği, video bilgisi çıkarma ve dosya yönetimi sağladım.')
add_bullet('Docker ile konteynerize ederek taşınabilir deployment mimarisi kurdum.')

# ═══ GİRİŞİMCİLİK ═══
add_section_title('Girişimcilik Deneyimi')

p = doc.add_paragraph()
run = p.add_run('NGD — New Generation Design · Kurucu & Ürün Geliştiricisi')
run.bold = True; run.font.size = Pt(10.5); run.font.name = 'Segoe UI'
run2 = p.add_run('\t2016 – 2022')
run2.font.size = Pt(9.5); run2.font.color.rgb = RGBColor(0x55,0x55,0x55)
p.paragraph_format.space_after = Pt(1)

add_simple_para('Termal yönetim, teknik tekstil ve savunma sanayii — Ar-Ge odaklı ürün geliştirme', 10, (0x33,0x33,0x33), space_after=2)

for b in [
    'Şirketi sıfırdan kurdu; müşteri problemini tanımladı, çözüm hipotezi oluşturdu, prototip geliştirdi ve MVP\'yi piyasaya çıkardı — modern bir start-up\'ın tüm aşamalarını bizzat yönetti.',
    'Müşteri ihtiyaç analizi, ürün-pazar uyumu araştırması ve ürün doğrulama süreçlerini yürüttü; B2B müşteriler için özel çözümler geliştirdi.',
    'Kurumsal web sitesi, dijital pazarlama, ürün lansmanı ve e-ticaret süreçlerini uçtan uca yönetti.',
    'Savunma sanayii için özel termal sistemler geliştirdi; Milli Savunma Bakanlığı ve Silahlı Kuvvetler\'e teknik sunum yaptı.',
    'Almanya, Fransa, İsveç, ABD ve Rusya dahil uluslararası tedarikçi ve ihracat süreçlerini koordine etti.',
]:
    add_bullet(b)

# ═══ MESLEKİ GEÇMİŞ ═══
add_section_title('Mesleki Geçmiş')

add_simple_para('Yurt içi ve yurt dışında yaklaşık 20 yıl boyunca teknik operasyon, proje koordinasyonu, ekip yönetimi ve kalite kontrol alanlarında görev aldı. Bu süre zarfında:', 10, (0x33,0x33,0x33), space_after=2)

for b in [
    'Farklı ülkelerde ve disiplinlerde saha deneyimi kazanarak değişen şartlara uyum sağlama ve sonuç odaklı hareket etme becerilerini geliştirdi.',
    'Teknik operasyon ve proje yönetiminde yüksek sorumluluk gerektiren görevlerde analitik bakış açısı ve problem çözme yetkinliğini ortaya koydu.',
    'Uluslararası projelerde müşteri ve tedarikçi iletişimi, kalite kontrol ve ekip koordinasyonu süreçlerini yürüttü.',
]:
    add_bullet(b)

p = doc.add_paragraph()
run = p.add_run('Bu deneyimler, yazılım geliştirme sürecinde ')
run.font.size = Pt(9.8); run.font.name = 'Segoe UI'
run = p.add_run('Agile thinking, technical problem solving, cross-functional collaboration ve project management')
run.bold = True; run.font.size = Pt(9.8); run.font.name = 'Segoe UI'
run = p.add_run(' gibi transferable skills olarak doğrudan katkı sağlamaktadır.')
run.font.size = Pt(9.8); run.font.name = 'Segoe UI'

# ═══ NİTELİKLER ═══
add_section_title('Nitelikler')
add_simple_para('Analitik düşünme · Problem çözme · Girişimci bakış açısı · Takım çalışmasına yatkın · Sonuç odaklı · Hızlı öğrenme', 9.8, (0x33,0x33,0x33), space_after=6)

# ═══ YABANCI DİL ═══
add_section_title('Yabancı Dil')

p = doc.add_paragraph()
run = p.add_run('Türkçe')
run.bold = True; run.font.size = Pt(10)
run2 = p.add_run(' — Ana dil')
run2.font.size = Pt(9.5)
p.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
run = p.add_run('İngilizce')
run.bold = True; run.font.size = Pt(10)
run2 = p.add_run(' — B2 / Upper-Intermediate · Teknik doküman okuma, yazışma ve temel konuşma')
run2.font.size = Pt(9.5)
p.paragraph_format.space_after = Pt(6)

# ═══ KARİYER HEDEFİ ═══
add_section_title('Kariyer Hedefi')
add_simple_para('Yapay zeka, LLM uygulamaları ve otomasyon odaklı yazılım geliştirme rollerinde kariyerini ilerletmeyi hedefleyen bir geliştirici. Mevcut full-stack proje deneyimini, AI/ML lisans eğitimiyle birleştirerek Junior AI Developer, Junior Backend Developer veya AI Automation Developer pozisyonlarında değer yaratmaya hazır.', 10, (0x22,0x22,0x22), space_after=6)

# ═══ REFERANSLAR ═══
add_section_title('Referanslar')
add_simple_para('Talep edilmesi halinde profesyonel referanslarım paylaşılacaktır.', 9.8, (0x55,0x55,0x55))

# Save
doc.save('cengiz_kilic_cv_tr.docx')
print(f"DOCX generated! Size: {os.path.getsize('cengiz_kilic_cv_tr.docx')} bytes")
